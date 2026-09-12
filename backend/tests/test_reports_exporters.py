"""Tests for app/reports/exporters.py - Phase 2's real PDF/DOCX report
exports. Both formats are rendered from the exact same markdown
generate_report_markdown() already produces, so these tests focus on:
the export actually being a valid file of the right type, the report's
facts surviving the markdown -> HTML -> PDF/DOCX conversion, and the
filename sanitizer never letting a project name inject a header/path.
"""

import io
import uuid

import pymupdf
from docx import Document

from app.models.case_file import CaseFile
from app.reports.exporters import render_docx, render_pdf, safe_report_filename


def _case_file(project_name: str = "Test Warehouse") -> CaseFile:
    case_file = CaseFile(session_id=str(uuid.uuid4()), project_name=project_name, state="Gujarat")
    case_file.classification_result.table_7_ref = "7H"
    return case_file


def test_render_pdf_produces_a_valid_pdf_with_report_facts():
    pdf_bytes = render_pdf(_case_file())

    assert pdf_bytes.startswith(b"%PDF")
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    text = "\n".join(page.get_text() for page in doc)
    assert "Test Warehouse" in text
    assert "7H" in text
    assert "advisory only" in text


def test_render_docx_produces_a_valid_docx_with_report_facts():
    docx_bytes = render_docx(_case_file())

    assert docx_bytes.startswith(b"PK")
    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Warehouse" in full_text
    assert "7H" in full_text
    # The report's own title is an H1 - confirm it round-tripped as a
    # heading, not just plain text, so the export is a real structured
    # document rather than one long paragraph.
    assert doc.paragraphs[0].style.name == "Heading 1"


def test_safe_report_filename_strips_header_and_path_unsafe_characters():
    unsafe = _case_file(project_name='../../etc/passwd"; evil="\r\nX-Injected: yes')

    result = safe_report_filename(unsafe)

    assert result == "etcpasswd evilX-Injected yes"
    for char in ('"', "\r", "\n", "/", ";", ":", "="):
        assert char not in result


def test_safe_report_filename_falls_back_when_project_name_is_empty():
    assert safe_report_filename(_case_file(project_name="")) == "fire-safety-report"
